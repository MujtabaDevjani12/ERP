"""
Excel import framework: Upload -> Validate -> Preview -> Show errors -> Confirm -> Import.

    job = create_import_job(company, "customers", uploaded_file, user)   # parses + validates
    job.status        # 'validated' (ready) or 'failed' (see job.errors); job.rows = preview rows
    confirm_import(job, user)                                           # only works when validated

Never imports invalid data: one bad row blocks the whole file, and every
problem is listed with its Excel row number.

Other apps add types with register_import_type() (e.g. hr registers 'attendance').
"""
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from typing import Callable

from django.db import transaction

from audit.models import AuditLog
from audit.services import log_action
from core.errors import ERPError
from core.money import D, q2
from .models import CustomerProfile, ImportJob, ItemProfile, PaymentTerm

MAX_ROWS = 20000


@dataclass
class ImportType:
    name: str
    label: str
    required: tuple
    optional: tuple
    validate_row: Callable      # (row: dict, ctx: dict) -> (clean: dict, errors: list[str])
    commit: Callable            # (clean_rows: list[dict], ctx: dict) -> None
    prepare: Callable = None    # optional: (ctx) -> None, builds lookups once


IMPORT_TYPES = {}


def register_import_type(name, *, label, required, optional=(), validate_row, commit, prepare=None):
    IMPORT_TYPES[name] = ImportType(name, label, tuple(required), tuple(optional), validate_row, commit, prepare)


# ---------------------------------------------------------------- parsing helpers
def parse_bool(value, default=False):
    if value in (None, ""):
        return default
    return str(value).strip().lower() in ("1", "true", "yes", "y")


def parse_decimal(value, field, errors, *, default=None, minimum=None):
    if value in (None, ""):
        return default
    try:
        number = Decimal(str(value).replace(",", "").strip())
    except InvalidOperation:
        errors.append(f"{field}: '{value}' is not a number")
        return default
    if minimum is not None and number < minimum:
        errors.append(f"{field}: must be at least {minimum}")
    return number


def read_workbook(file):
    from openpyxl import load_workbook
    sheet = load_workbook(file, read_only=True, data_only=True).active
    rows = sheet.iter_rows(values_only=True)
    try:
        header = [str(h or "").strip().lower().replace(" ", "_") for h in next(rows)]
    except StopIteration:
        raise ERPError("The file is empty.")
    out = []
    for number, raw in enumerate(rows, start=2):
        if raw is None or all(v in (None, "") for v in raw):
            continue
        out.append({"_row": number, **{h: (raw[i] if i < len(raw) else None) for i, h in enumerate(header) if h}})
        if len(out) > MAX_ROWS:
            raise ERPError(f"Too many rows (max {MAX_ROWS}). Split the file.")
    return header, out


def _jsonable(value):
    if isinstance(value, Decimal):
        return str(value)
    if hasattr(value, "isoformat"):
        return value.isoformat()
    return value


# ---------------------------------------------------------------- job lifecycle
def create_import_job(company, import_type, file, user=None) -> ImportJob:
    spec = IMPORT_TYPES.get(import_type)
    if spec is None:
        raise ERPError(f"Unknown import type '{import_type}'. Available: {', '.join(sorted(IMPORT_TYPES))}")
    job = ImportJob.objects.create(company=company, import_type=import_type,
                                   uploaded_by=user if getattr(user, "is_authenticated", False) else None)
    try:
        header, rows = read_workbook(file)
    except ERPError as exc:
        job.status, job.errors, job.summary = ImportJob.FAILED, [{"row": 1, "errors": [str(exc)]}], str(exc)
        job.save()
        return job
    except Exception as exc:                                  # unreadable / not an xlsx
        job.status, job.errors, job.summary = ImportJob.FAILED, [{"row": 1, "errors": [f"Could not read the file: {exc}"]}], "Unreadable file"
        job.save()
        return job
    missing = [c for c in spec.required if c not in header]
    if missing:
        job.status = ImportJob.FAILED
        job.errors = [{"row": 1, "errors": [f"Missing required column(s): {', '.join(missing)}"]}]
        job.summary = "Missing columns"
        job.save()
        return job
    return validate_rows(job, spec, rows)


def validate_rows(job, spec, rows):
    ctx = {"company": job.company, "seen": set()}
    if spec.prepare:
        spec.prepare(ctx)
    clean_rows, errors = [], []
    for row in rows:
        clean, row_errors = spec.validate_row(row, ctx)
        if row_errors:
            errors.append({"row": row["_row"], "errors": row_errors})
        else:
            clean["_row"] = row["_row"]
            clean_rows.append({k: _jsonable(v) for k, v in clean.items()})
    job.rows = clean_rows
    job.errors = errors
    if errors:
        job.status = ImportJob.FAILED
        job.summary = f"{len(errors)} row(s) have errors; nothing will be imported until they are fixed."
    elif not clean_rows:
        job.status = ImportJob.FAILED
        job.summary = "The file has no data rows."
        job.errors = [{"row": 1, "errors": ["No data rows found."]}]
    else:
        job.status = ImportJob.VALIDATED
        job.summary = f"{len(clean_rows)} row(s) are valid and ready to import."
    job.save()
    return job


@transaction.atomic
def confirm_import(job, user=None) -> ImportJob:
    job = ImportJob.objects.select_for_update().get(pk=job.pk)
    if job.status != ImportJob.VALIDATED:
        raise ERPError("Only a validated import (with no errors) can be confirmed.")
    spec = IMPORT_TYPES[job.import_type]
    ctx = {"company": job.company, "user": user, "seen": set()}
    if spec.prepare:
        spec.prepare(ctx)
    spec.commit(job.rows, ctx)
    job.status = ImportJob.IMPORTED
    job.summary = f"{len(job.rows)} row(s) imported."
    job.save(update_fields=["status", "summary", "updated_at"])
    log_action(AuditLog.CREATED, job, user=user, note=f"imported {len(job.rows)} {job.import_type} row(s)", company=job.company)
    return job


# ---------------------------------------------------------------- shared lookups
def _stores(ctx):
    from tenants.models import Store
    stores = list(Store.objects.filter(company=ctx["company"]).order_by("id"))
    ctx["stores"] = {s.name.lower(): s for s in stores}
    ctx["default_store"] = stores[0] if stores else None


def _store_for(row, ctx, errors):
    name = str(row.get("store") or "").strip().lower()
    if name:
        store = ctx["stores"].get(name)
        if store is None:
            errors.append(f"store: '{row.get('store')}' not found")
        return store
    if ctx["default_store"] is None:
        errors.append("this company has no branch (store)")
    return ctx["default_store"]


def _payment_term(row, ctx, errors, key="payment_term"):
    name = str(row.get(key) or "").strip()
    if not name:
        return None
    term = PaymentTerm.objects.filter(company=ctx["company"], name__iexact=name).first()
    if term is None:
        errors.append(f"{key}: '{name}' is not defined (Master data > Payment terms)")
    return term


def _str(row, key, limit=255):
    return str(row.get(key) or "").strip()[:limit]


# ---------------------------------------------------------------- customers
def _validate_customer(row, ctx):
    errors = []
    name = _str(row, "name", 150)
    if not name:
        errors.append("name is required")
    store = _store_for(row, ctx, errors)
    ctype = (_str(row, "customer_type") or "retail").lower()
    if ctype not in dict(CustomerProfile.TYPE_CHOICES):
        errors.append(f"customer_type: '{ctype}' must be one of {', '.join(dict(CustomerProfile.TYPE_CHOICES))}")
    credit = parse_decimal(row.get("credit_limit"), "credit_limit", errors, default=Decimal("0"), minimum=Decimal("0"))
    opening = parse_decimal(row.get("opening_balance"), "opening_balance", errors, default=Decimal("0"))
    term = _payment_term(row, ctx, errors)
    key = ("cust", name.lower(), _str(row, "phone").lower())
    if key in ctx["seen"]:
        errors.append("duplicate customer inside this file")
    ctx["seen"].add(key)
    if store is not None and name:
        from customers.models import Customer
        if Customer.objects.filter(store=store, name__iexact=name, phone=_str(row, "phone", 30)).exists():
            errors.append("customer already exists")
    clean = {
        "name": name, "store": store.name if store else "", "phone": _str(row, "phone", 30), "email": _str(row, "email"),
        "cnic": _str(row, "cnic", 15), "address": _str(row, "address"), "customer_type": ctype,
        "company_name": _str(row, "company_name", 150), "city": _str(row, "city", 80), "tax_number": _str(row, "tax_number", 60),
        "credit_limit": credit, "opening_balance": opening, "payment_term": term.name if term else "",
    }
    return clean, errors


def _commit_customers(rows, ctx):
    from customers.models import Customer
    from .services import ensure_customer_profile
    company = ctx["company"]
    for r in rows:
        store = ctx["stores"][r["store"].lower()]
        customer = Customer.objects.create(store=store, name=r["name"], phone=r["phone"], email=r["email"],
                                           cnic=r["cnic"], address=r["address"])
        term = PaymentTerm.objects.filter(company=company, name__iexact=r["payment_term"]).first() if r["payment_term"] else None
        ensure_customer_profile(customer, company, customer_type=r["customer_type"], company_name=r["company_name"],
                                city=r["city"], tax_number=r["tax_number"], credit_limit=D(r["credit_limit"]),
                                opening_balance=D(r["opening_balance"]), payment_term=term)


register_import_type("customers", label="Customers", required=("name",),
                     optional=("store", "phone", "email", "cnic", "address", "customer_type", "company_name", "city",
                               "tax_number", "credit_limit", "opening_balance", "payment_term"),
                     validate_row=_validate_customer, commit=_commit_customers, prepare=_stores)


# ---------------------------------------------------------------- vendors
def _validate_vendor(row, ctx):
    errors = []
    name = _str(row, "name", 150)
    if not name:
        errors.append("name is required")
    store = _store_for(row, ctx, errors)
    credit = parse_decimal(row.get("credit_limit"), "credit_limit", errors, default=Decimal("0"), minimum=Decimal("0"))
    opening = parse_decimal(row.get("opening_balance"), "opening_balance", errors, default=Decimal("0"))
    term = _payment_term(row, ctx, errors)
    key = ("vend", name.lower())
    if key in ctx["seen"]:
        errors.append("duplicate vendor inside this file")
    ctx["seen"].add(key)
    if store is not None and name:
        from suppliers.models import Supplier
        if Supplier.objects.filter(store=store, name__iexact=name).exists():
            errors.append("vendor already exists")
    clean = {"name": name, "store": store.name if store else "", "contact_person": _str(row, "contact_person", 100),
             "phone": _str(row, "phone", 30), "email": _str(row, "email"), "address": _str(row, "address"),
             "city": _str(row, "city", 80), "tax_number": _str(row, "tax_number", 60), "credit_limit": credit,
             "opening_balance": opening, "payment_term": term.name if term else ""}
    return clean, errors


def _commit_vendors(rows, ctx):
    from suppliers.models import Supplier
    from .services import ensure_vendor_profile
    company = ctx["company"]
    for r in rows:
        store = ctx["stores"][r["store"].lower()]
        supplier = Supplier.objects.create(store=store, name=r["name"], contact_person=r["contact_person"],
                                           phone=r["phone"], email=r["email"], address=r["address"])
        term = PaymentTerm.objects.filter(company=company, name__iexact=r["payment_term"]).first() if r["payment_term"] else None
        ensure_vendor_profile(supplier, company, city=r["city"], tax_number=r["tax_number"],
                              credit_limit=D(r["credit_limit"]), opening_balance=D(r["opening_balance"]), payment_term=term)


register_import_type("vendors", label="Vendors", required=("name",),
                     optional=("store", "contact_person", "phone", "email", "address", "city", "tax_number",
                               "credit_limit", "opening_balance", "payment_term"),
                     validate_row=_validate_vendor, commit=_commit_vendors, prepare=_stores)


# ---------------------------------------------------------------- items
def _validate_item(row, ctx):
    from catalog.models import Product, ProductVariant, Unit
    errors = []
    sku, name = _str(row, "sku", 60), _str(row, "name", 180)
    if not sku:
        errors.append("sku is required")
    if not name:
        errors.append("name is required")
    store = _store_for(row, ctx, errors)
    unit = None
    if store is not None:
        unit = Unit.objects.filter(store=store, short_name__iexact=_str(row, "unit", 10)).first() or \
               Unit.objects.filter(store=store, name__iexact=_str(row, "unit", 50)).first()
        if unit is None:
            errors.append(f"unit: '{row.get('unit')}' not found in that branch (Products > Units)")
    price = parse_decimal(row.get("selling_price"), "selling_price", errors, default=Decimal("0"), minimum=Decimal("0"))
    cost = parse_decimal(row.get("cost_price"), "cost_price", errors, default=Decimal("0"), minimum=Decimal("0"))
    std = parse_decimal(row.get("standard_cost"), "standard_cost", errors, default=Decimal("0"), minimum=Decimal("0"))
    reorder = parse_decimal(row.get("reorder_level"), "reorder_level", errors, default=Decimal("0"), minimum=Decimal("0"))
    item_type = (_str(row, "item_type") or ItemProfile.TRADING).lower()
    if item_type not in dict(ItemProfile.TYPE_CHOICES):
        errors.append(f"item_type: must be one of {', '.join(dict(ItemProfile.TYPE_CHOICES))}")
    barcode = _str(row, "barcode", 64)
    if barcode and (ProductVariant.objects.filter(barcode=barcode).exists() or ("barcode", barcode) in ctx["seen"]):
        errors.append(f"barcode '{barcode}' already exists")
    if barcode:
        ctx["seen"].add(("barcode", barcode))
    if sku:
        if ("sku", sku.lower()) in ctx["seen"]:
            errors.append("duplicate sku inside this file")
        ctx["seen"].add(("sku", sku.lower()))
        if store is not None and Product.objects.filter(store=store, sku=sku).exists():
            errors.append(f"sku '{sku}' already exists")
        if ProductVariant.objects.filter(sku=f"{sku}-DEF").exists():
            errors.append(f"variant sku '{sku}-DEF' already exists")
    clean = {"sku": sku, "name": name, "store": store.name if store else "", "unit_id": unit.id if unit else None,
             "selling_price": price, "cost_price": cost, "standard_cost": std, "reorder_level": reorder,
             "barcode": barcode, "category": _str(row, "category", 100), "brand": _str(row, "brand", 100),
             "item_type": item_type, "batch_required": parse_bool(row.get("batch_required")),
             "expiry_required": parse_bool(row.get("expiry_required")), "qc_required": parse_bool(row.get("qc_required")),
             "color": _str(row, "color", 40), "size": _str(row, "size", 40), "gsm": _str(row, "gsm", 20),
             "composition": _str(row, "composition", 120)}
    return clean, errors


def _commit_items(rows, ctx):
    from catalog.models import Brand, Category, Product, Unit
    for r in rows:
        store = ctx["stores"][r["store"].lower()]
        category = Category.objects.get_or_create(store=store, parent=None, name=r["category"])[0] if r["category"] else None
        brand = Brand.objects.get_or_create(store=store, name=r["brand"])[0] if r["brand"] else None
        product = Product.objects.create(store=store, category=category, brand=brand, unit=Unit.objects.get(pk=r["unit_id"]),
                                         name=r["name"], sku=r["sku"], product_type="single")
        variant = product.variants.first()          # created by catalog's post_save signal
        variant.selling_price, variant.cost_price = D(r["selling_price"]), D(r["cost_price"])
        if r["barcode"]:
            variant.barcode = r["barcode"]
        variant.save()
        gsm = D(r["gsm"]) if r["gsm"] else None
        ItemProfile.objects.create(
            variant=variant, item_type=r["item_type"], batch_required=r["batch_required"],
            expiry_required=r["expiry_required"], qc_required_on_receipt=r["qc_required"], color=r["color"], size=r["size"],
            gsm=gsm, composition=r["composition"], reorder_level=D(r["reorder_level"]), standard_cost=D(r["standard_cost"]),
        )


register_import_type("items", label="Items", required=("sku", "name", "unit"),
                     optional=("store", "selling_price", "cost_price", "standard_cost", "reorder_level", "barcode", "category",
                               "brand", "item_type", "batch_required", "expiry_required", "qc_required", "color", "size",
                               "gsm", "composition"),
                     validate_row=_validate_item, commit=_commit_items, prepare=_stores)


# ---------------------------------------------------------------- prices
def _validate_price(row, ctx):
    from catalog.models import ProductVariant
    errors = []
    sku = _str(row, "sku", 80)
    variant = ProductVariant.objects.filter(sku=sku, product__store__company=ctx["company"]).first() if sku else None
    if variant is None:
        errors.append(f"sku '{sku}' not found")
    price = parse_decimal(row.get("selling_price"), "selling_price", errors, default=None, minimum=Decimal("0"))
    cost = parse_decimal(row.get("cost_price"), "cost_price", errors, default=None, minimum=Decimal("0"))
    if price is None and cost is None:
        errors.append("give selling_price and/or cost_price")
    return {"sku": sku, "variant_id": variant.id if variant else None, "selling_price": price, "cost_price": cost}, errors


def _commit_prices(rows, ctx):
    from catalog.models import ProductVariant
    for r in rows:
        variant = ProductVariant.objects.get(pk=r["variant_id"])
        if r["selling_price"] is not None:
            variant.selling_price = D(r["selling_price"])
        if r["cost_price"] is not None:
            variant.cost_price = D(r["cost_price"])
        variant.save()


register_import_type("prices", label="Prices", required=("sku",), optional=("selling_price", "cost_price"),
                     validate_row=_validate_price, commit=_commit_prices)


# ---------------------------------------------------------------- opening stock
def _prepare_warehouses(ctx):
    from tenants.models import Warehouse
    whs = Warehouse.objects.filter(store__company=ctx["company"]).select_related("store")
    ctx["warehouses"] = {}
    for w in whs:
        ctx["warehouses"][w.name.lower()] = w
        if w.code:
            ctx["warehouses"][w.code.lower()] = w


def _validate_opening_stock(row, ctx):
    from catalog.models import ProductVariant
    errors = []
    sku = _str(row, "sku", 80)
    variant = ProductVariant.objects.filter(sku=sku, product__store__company=ctx["company"]).first() if sku else None
    if variant is None:
        errors.append(f"sku '{sku}' not found")
    wh = ctx["warehouses"].get(_str(row, "warehouse", 120).lower())
    if wh is None:
        errors.append(f"warehouse '{row.get('warehouse')}' not found (use name or code)")
    qty = parse_decimal(row.get("quantity"), "quantity", errors, default=None)
    if qty is None or qty <= 0:
        errors.append("quantity must be greater than zero")
    cost = parse_decimal(row.get("unit_cost"), "unit_cost", errors, default=None, minimum=Decimal("0"))
    if cost is None:
        errors.append("unit_cost is required")
    key = ("os", sku, wh.id if wh else None)
    if key in ctx["seen"]:
        errors.append("same sku/warehouse listed twice")
    ctx["seen"].add(key)
    return {"sku": sku, "variant_id": variant.id if variant else None, "warehouse_id": wh.id if wh else None,
            "quantity": qty, "unit_cost": cost}, errors


def _commit_opening_stock(rows, ctx):
    from accounting.services import create_journal_entry, get_account
    from catalog.models import ProductVariant
    from company.utils import company_of_warehouse
    from stock.services import inventory_account_key, post_stock_movement
    from tenants.models import Warehouse
    by_wh = {}
    for r in rows:
        wh = Warehouse.objects.select_related("store").get(pk=r["warehouse_id"])
        variant = ProductVariant.objects.get(pk=r["variant_id"])
        qty, cost = D(r["quantity"]), D(r["unit_cost"])
        post_stock_movement(variant, wh, qty, reference_type="OPENING_STOCK", reference_no=f"IMPORT-{ctx['company'].pk}",
                            unit_cost=cost, user=ctx.get("user"), remarks="Opening stock import")
        by_wh.setdefault(wh, D(0))
        by_wh[wh] += qty * cost
    for wh, value in by_wh.items():
        company = company_of_warehouse(wh)
        create_journal_entry(
            company, voucher_type="OPENING", narration=f"Opening stock {wh.name}", user=ctx.get("user"), store=wh.store,
            lines=[{"account": get_account(company, inventory_account_key(wh)), "debit": q2(value)},
                   {"account": get_account(company, "opening_balance_equity"), "credit": q2(value)}],
        )


register_import_type("opening_stock", label="Opening stock", required=("sku", "warehouse", "quantity", "unit_cost"),
                     validate_row=_validate_opening_stock, commit=_commit_opening_stock, prepare=_prepare_warehouses)


# ---------------------------------------------------------------- opening balances (journal)
def _prepare_accounts(ctx):
    from accounting.models import Account
    ctx["accounts"] = {a.code: a for a in Account.objects.filter(company=ctx["company"])}


def _validate_opening_balance(row, ctx):
    errors = []
    code = _str(row, "account_code", 20)
    account = ctx["accounts"].get(code)
    if account is None:
        errors.append(f"account_code '{code}' not found")
    elif account.is_group:
        errors.append(f"account '{code}' is a group account")
    debit = parse_decimal(row.get("debit"), "debit", errors, default=Decimal("0"), minimum=Decimal("0"))
    credit = parse_decimal(row.get("credit"), "credit", errors, default=Decimal("0"), minimum=Decimal("0"))
    if debit > 0 and credit > 0:
        errors.append("a line cannot have both debit and credit")
    if debit == 0 and credit == 0:
        errors.append("debit or credit is required")
    totals = ctx.setdefault("totals", [Decimal("0"), Decimal("0")])
    totals[0] += debit or Decimal("0")
    totals[1] += credit or Decimal("0")
    return {"account_code": code, "debit": debit, "credit": credit, "description": _str(row, "description")}, errors


def _commit_opening_balances(rows, ctx):
    from accounting.services import create_journal_entry
    from accounting.models import Account
    lines = []
    for r in rows:
        account = Account.objects.get(company=ctx["company"], code=r["account_code"])
        lines.append({"account": account, "debit": D(r["debit"]), "credit": D(r["credit"]), "description": r["description"]})
    # create_journal_entry itself refuses an unbalanced set, so the import fails as a whole
    create_journal_entry(ctx["company"], lines=lines, voucher_type="OPENING", narration="Opening balances (Excel import)",
                         user=ctx.get("user"))


register_import_type("opening_balances", label="Opening balances", required=("account_code",),
                     optional=("debit", "credit", "description"),
                     validate_row=_validate_opening_balance, commit=_commit_opening_balances, prepare=_prepare_accounts)
