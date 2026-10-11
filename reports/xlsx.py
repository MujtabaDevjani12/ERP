"""Excel (.xlsx) and CSV reading/writing for report exports and imports.

``openpyxl`` is imported lazily so the rest of the app keeps working if it is not installed yet; the Excel features then
answer with a clear "install openpyxl" message instead of crashing the server.
"""
from __future__ import annotations

import csv
import io
import re
from datetime import date, datetime
from decimal import Decimal

from django.http import HttpResponse
from django.utils import timezone


XLSX_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_IMPORT_ROWS = 5000


class ExcelError(Exception):
    """A problem the user can fix (missing library, wrong file...). The message is safe to show as it is."""


def _openpyxl():
    try:
        import openpyxl  # noqa: F401
        return openpyxl
    except ImportError:  # pragma: no cover - depends on the environment
        raise ExcelError("Excel support is not installed on the server. Run: pip install openpyxl")


def safe_cell(value):
    """Make a value safe for a spreadsheet cell: no formulas can be smuggled in through text (``=cmd...``), Decimals become numbers,
    aware datetimes become local time without a timezone (Excel cannot store one)."""
    if value is None:
        return None
    if isinstance(value, bool):
        return "Yes" if value else "No"
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, datetime):
        if timezone.is_aware(value):
            value = timezone.localtime(value).replace(tzinfo=None)
        return value
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        return value
    text = str(value)
    if text[:1] in ("=", "+", "-", "@", "\t", "\r"):
        return "'" + text
    return text


def build_workbook(sheets: list[dict]) -> bytes:
    """``sheets`` = [{"title", "headers": [...], "rows": [[...]], "summary": [(label, value)], "heading": str, "notes": [str], "money_cols": [idx]}]."""
    openpyxl = _openpyxl()
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    head_fill = PatternFill("solid", fgColor="1F2937")
    for sh in sheets:
        ws = wb.create_sheet(re.sub(r"[\[\]\*\?/\\:]", "-", str(sh.get("title") or "Sheet"))[:31])
        r = 1
        if sh.get("heading"):
            ws.cell(r, 1, safe_cell(sh["heading"])).font = Font(bold=True, size=14)
            r += 1
        for note in sh.get("notes") or []:
            ws.cell(r, 1, safe_cell(note)).font = Font(italic=True, color="6B7280")
            r += 1
        if sh.get("summary"):
            r += 1
            for label, value in sh["summary"]:
                ws.cell(r, 1, safe_cell(label)).font = Font(bold=True)
                c = ws.cell(r, 2, safe_cell(value))
                if isinstance(c.value, float):
                    c.number_format = "#,##0.00"
                c.alignment = Alignment(horizontal="right")
                r += 1
        if sh.get("heading") or sh.get("notes") or sh.get("summary"):
            r += 1
        headers = sh.get("headers") or []
        first = r
        for i, h in enumerate(headers, 1):
            c = ws.cell(r, i, safe_cell(h))
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = head_fill
            c.alignment = Alignment(vertical="center", wrap_text=True)
        widths = [max(10, min(len(str(h)) + 2, 40)) for h in headers]
        for row in sh.get("rows") or []:
            r += 1
            for i, v in enumerate(row, 1):
                cell = ws.cell(r, i, safe_cell(v))
                if isinstance(cell.value, float):
                    cell.number_format = "#,##0.00" if not float(cell.value).is_integer() or i - 1 in (sh.get("money_cols") or ()) else "#,##0"
                elif isinstance(cell.value, datetime):
                    cell.number_format = "yyyy-mm-dd hh:mm"
                elif isinstance(cell.value, date):
                    cell.number_format = "yyyy-mm-dd"
                if i <= len(widths):
                    widths[i - 1] = max(widths[i - 1], min(len(str(cell.value if cell.value is not None else "")) + 2, 48))
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[get_column_letter(i)].width = w
        if headers:
            ws.freeze_panes = ws.cell(first + 1, 1)
            if sh.get("rows"):
                ws.auto_filter.ref = f"A{first}:{get_column_letter(len(headers))}{r}"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def xlsx_response(filename: str, sheets: list[dict]) -> HttpResponse:
    resp = HttpResponse(build_workbook(sheets), content_type=XLSX_TYPE)
    safe = re.sub(r"[^A-Za-z0-9._-]+", "-", filename).strip("-") or "export"
    resp["Content-Disposition"] = f'attachment; filename="{safe}.xlsx"'
    resp["Cache-Control"] = "no-store"
    return resp


def norm_header(h) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(h or "").strip().lower()).strip("_")


def read_table(upload, aliases: dict[str, str] | None = None) -> tuple[list[str], list[tuple[int, dict]]]:
    """Read the first sheet of an .xlsx (or a .csv) upload. Returns (headers, [(row_number, {header: value})]). Blank rows are skipped.
    Headers are normalised (``Unit cost`` -> ``unit_cost``) and mapped through ``aliases``."""
    if upload is None:
        raise ExcelError("Choose an Excel (.xlsx) or CSV file first.")
    if upload.size > MAX_UPLOAD_BYTES:
        raise ExcelError("That file is larger than 5 MB. Split it into smaller files.")
    raw = upload.read()
    name = (getattr(upload, "name", "") or "").lower()
    aliases = aliases or {}
    rows: list[list] = []
    if raw[:2] == b"PK":  # a zip container = .xlsx
        openpyxl = _openpyxl()
        try:
            wb = openpyxl.load_workbook(io.BytesIO(raw), read_only=True, data_only=True)
            ws = wb.worksheets[0]
            for i, r in enumerate(ws.iter_rows(values_only=True)):
                rows.append(list(r))
                if i > MAX_IMPORT_ROWS + 50:
                    break
        except Exception:
            raise ExcelError("This file could not be read as an Excel workbook. Save it as .xlsx and try again.")
    elif name.endswith((".csv", ".txt")):
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            text = raw.decode("latin-1")
        try:
            dialect = csv.Sniffer().sniff(text[:2048], delimiters=",;\t")
        except csv.Error:
            dialect = csv.excel
        rows = [list(r) for r in csv.reader(io.StringIO(text), dialect)]
    else:
        raise ExcelError("Upload an Excel (.xlsx) or CSV file.")
    # the header row is the first row that has at least one non-empty cell
    start = next((i for i, r in enumerate(rows) if any(c not in (None, "") for c in r)), None)
    if start is None:
        raise ExcelError("The file is empty.")
    headers = [aliases.get(norm_header(h), norm_header(h)) for h in rows[start]]
    out = []
    for i, r in enumerate(rows[start + 1:], start + 2):
        if not any(c not in (None, "") for c in r):
            continue
        out.append((i, {h: (r[j] if j < len(r) else None) for j, h in enumerate(headers) if h}))
    if len(out) > MAX_IMPORT_ROWS:
        raise ExcelError(f"At most {MAX_IMPORT_ROWS} rows can be imported at once. Split the file.")
    return [h for h in headers if h], out


# ------------------------------------------------------------------ cell parsing helpers used by the importers
def text(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)  # a code typed as 123 must not become "123.0"
    return str(v).strip()


def number(v, field="Value", default=None, minimum=None) -> Decimal | None:
    if v is None or str(v).strip() == "":
        return default
    try:
        n = Decimal(str(v).replace(",", "").strip())
    except Exception:
        raise ValueError(f"{field} must be a number (got “{v}”).")
    if not n.is_finite():
        raise ValueError(f"{field} must be a number.")
    if minimum is not None and n < minimum:
        raise ValueError(f"{field} cannot be less than {minimum}.")
    return n


def yes_no(v, default=None) -> bool | None:
    if v is None or str(v).strip() == "":
        return default
    s = str(v).strip().lower()
    if s in ("1", "y", "yes", "true", "t", "active", "on"):
        return True
    if s in ("0", "n", "no", "false", "f", "inactive", "off"):
        return False
    raise ValueError(f"“{v}” is not Yes or No.")


def when(v, field="Date") -> datetime | None:
    if v is None or str(v).strip() == "":
        return None
    if isinstance(v, datetime):
        d = v
    elif isinstance(v, date):
        d = datetime.combine(v, datetime.min.time())
    else:
        s = str(v).strip()
        for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d %H:%M", "%Y-%m-%d", "%d/%m/%Y %H:%M", "%d/%m/%Y", "%d-%m-%Y", "%d-%b-%Y"):
            try:
                d = datetime.strptime(s, fmt)
                break
            except ValueError:
                continue
        else:
            raise ValueError(f"{field} “{s}” is not a date. Use YYYY-MM-DD.")
    return timezone.make_aware(d) if timezone.is_naive(d) else d
