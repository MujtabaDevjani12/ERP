"""Excel import / export: templates, exports, dry run, row-by-row errors, matching rules and branch isolation."""
import io
from decimal import Decimal

import openpyxl
from django.core.files.uploadedfile import SimpleUploadedFile
from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from catalog.models import Brand, Category, Product, ProductVariant, Unit
from core.pos_fixtures import PASSWORD, PosWorld, grant
from customers.models import Customer
from expenses.models import Expense, ExpenseCategory
from inventory.models import StockItem
from suppliers.models import PurchaseOrder, Supplier
from tenants.models import Store, Warehouse


def xlsx(headers, rows, name="data.xlsx"):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(headers)
    for r in rows:
        ws.append(r)
        for cell in ws[ws.max_row]:
            if isinstance(cell.value, str) and cell.value.startswith("="):
                cell.data_type = "s"                  # typed as text, like a cell the user wrote with a leading apostrophe
    buf = io.BytesIO()
    wb.save(buf)
    return SimpleUploadedFile(name, buf.getvalue(), content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")


def sheet_rows(response):
    ws = openpyxl.load_workbook(io.BytesIO(response.content)).worksheets[0]
    return [list(r) for r in ws.iter_rows(values_only=True)]


class ExchangeTests(TestCase):
    def setUp(self):
        self.w = PosWorld()
        self.client.login(username="manager", password=PASSWORD)
        self.today = timezone.localdate()

    def imp(self, key, headers, rows, mode="import"):
        r = self.client.post(reverse("reports:data_import", args=[key]), {"file": xlsx(headers, rows), "mode": mode})
        self.assertEqual(r.status_code, 200)
        return r.context["result"] if r.context.get("result") else r.context

    # ------------------------------------------------------------------ pages and downloads
    def test_hub_templates_and_exports_work_for_every_list(self):
        self.assertEqual(self.client.get(reverse("reports:data_hub")).status_code, 200)
        from reports.exchange import REGISTRY
        for key, res in REGISTRY.items():
            ex = self.client.get(reverse("reports:data_export", args=[key]))
            self.assertEqual(ex.status_code, 200, key)
            self.assertIn("spreadsheetml", ex["Content-Type"])
            self.assertTrue(sheet_rows(ex), key)
            if not res.export_only:
                tpl = self.client.get(reverse("reports:data_template", args=[key]))
                self.assertEqual(tpl.status_code, 200, key)
                wb = openpyxl.load_workbook(io.BytesIO(tpl.content))
                self.assertEqual(wb.sheetnames[-1], "How to fill")
            else:
                self.assertEqual(self.client.get(reverse("reports:data_template", args=[key])).status_code, 403)
                self.assertEqual(self.client.post(reverse("reports:data_import", args=[key]), {"file": xlsx(["a"], [[1]])}).status_code, 403)

    def test_unknown_list_is_404(self):
        self.assertEqual(self.client.get(reverse("reports:data_export", args=["nothing"])).status_code, 404)

    # ---------------------------------------------------------------- products
    def test_new_single_product_with_opening_stock(self):
        res = self.imp("products", ["SKU", "Name", "Category", "Brand", "Unit", "Cost price", "Selling price", "Opening stock"],
                       [["CHAI-1", "Chai 250g", "Grocery", "Lipton", "Packet", 40, 60, 24]])
        self.assertEqual((res["created"], res["failed"]), (1, 0))
        self.assertEqual(dict(res["made"]), {"brands": 1, "categories": 1, "units": 1})
        v = ProductVariant.objects.get(product__sku="CHAI-1")
        self.assertEqual((v.name, v.selling_price, v.cost_price, v.product.category.name, v.product.brand.name, v.product.product_type), ("Default", 60, 40, "Grocery", "Lipton", "single"))
        self.assertEqual(StockItem.objects.get(variant=v, warehouse=self.w.warehouse).quantity, 24)
        self.assertEqual(ProductVariant.objects.filter(product__sku="CHAI-1").count(), 1)          # no stray second variant

    def test_product_with_variants(self):
        res = self.imp("products", ["SKU", "Name", "Selling price", "Variants"], [["TEE", "T-shirt", 500, "Small:450; Large:550:300; Medium"]])
        self.assertEqual((res["created"], res["failed"]), (1, 0))
        p = Product.objects.get(sku="TEE")
        self.assertEqual(p.product_type, "variant")
        got = {v.name: (v.selling_price, v.cost_price) for v in p.variants.all()}
        self.assertEqual(got, {"Small": (450, 0), "Large": (550, 300), "Medium": (500, 0)})
        self.assertEqual(len({v.sku for v in p.variants.all()}), 3)
        # run again with a price change + a new size: existing variants are updated, nothing is removed
        res = self.imp("products", ["SKU", "Name", "Variants"], [["TEE", "", "Small:470; XL:600"]])
        self.assertEqual((res["updated"], res["failed"]), (1, 0))
        got = {v.name: v.selling_price for v in p.variants.all()}
        self.assertEqual(got, {"Small": 470, "Large": 550, "Medium": 500, "XL": 600})

    def test_update_existing_product_by_sku_or_variant_code(self):
        res = self.imp("products", ["SKU", "Selling price", "Cost price"], [["TEA", 120, 70], ["MILK-1", 55, ""]])
        self.assertEqual((res["updated"], res["failed"]), (2, 0))
        self.w.tea.refresh_from_db(); self.w.milk.refresh_from_db()
        self.assertEqual((self.w.tea.selling_price, self.w.tea.cost_price, self.w.milk.selling_price, self.w.milk.cost_price), (120, 70, 55, 10))

    def test_check_only_saves_nothing(self):
        res = self.imp("products", ["SKU", "Name", "Selling price"], [["NEW-1", "New thing", 10]], mode="check")
        self.assertEqual((res["dry_run"], res["created"]), (True, 1))
        self.assertFalse(Product.objects.filter(sku="NEW-1").exists())
        self.assertFalse(Category.objects.exists())

    def test_bad_rows_are_reported_and_good_rows_still_import(self):
        res = self.imp("products", ["SKU", "Name", "Selling price", "Barcode"],
                       [["OK-1", "Fine", 10, ""], ["BAD-1", "No price", "", ""], ["BAD-2", "Negative", -5, ""], ["BAD-3", "Dup barcode", 5, "BC" + "TEA"], ["OK-2", "Also fine", "1,250.50", ""]])
        self.assertEqual((res["created"], res["failed"]), (2, 3))
        self.assertEqual([e["row"] for e in res["errors"]], [3, 4, 5])
        self.assertIn("price", res["errors"][0]["message"].lower())
        self.assertIn("barcode", res["errors"][2]["message"].lower())
        self.assertEqual(Product.objects.get(sku="OK-2").variants.get().selling_price, Decimal("1250.50"))

    def test_failed_row_does_not_leave_half_a_product_or_phantom_category(self):
        res = self.imp("products", ["SKU", "Name", "Selling price", "Category", "Tax"], [["X-1", "One", 10, "Fresh", "No such tax"], ["X-2", "Two", 10, "Fresh", ""]])
        self.assertEqual((res["created"], res["failed"]), (1, 1))
        self.assertEqual(Category.objects.filter(name="Fresh").count(), 1)
        self.assertTrue(Product.objects.get(sku="X-2").category_id)
        self.assertFalse(Product.objects.filter(sku="X-1").exists())

    def test_export_can_be_imported_back_unchanged(self):
        ex = self.client.get(reverse("reports:data_export", args=["products"]))
        rows = sheet_rows(ex)
        file = xlsx(rows[0], rows[1:])
        res = self.client.post(reverse("reports:data_import", args=["products"]), {"file": file, "mode": "import"}).context["result"]
        self.assertEqual((res["failed"], res["created"]), (0, 0))
        self.assertEqual(res["updated"], 3)
        self.assertEqual(StockItem.objects.get(variant=self.w.tea, warehouse=self.w.warehouse).quantity, 10)      # stock untouched

    def test_formula_text_in_a_cell_is_exported_as_text(self):
        self.imp("products", ["SKU", "Name", "Selling price"], [["F-1", "=1+1", 5]])
        rows = sheet_rows(self.client.get(reverse("reports:data_export", args=["products"])))
        names = [r[1] for r in rows]
        self.assertIn("'=1+1", names)

    def test_missing_columns_are_a_clear_error(self):
        ctx = self.imp("products", ["Selling price"], [[1]])
        self.assertIn("missing the column", ctx["error"])
        ctx = self.imp("customers", ["Phone"], [["1"]])
        self.assertIn("Name", ctx["error"])

    def test_not_an_excel_file(self):
        r = self.client.post(reverse("reports:data_import", args=["products"]), {"file": SimpleUploadedFile("x.pdf", b"%PDF-1.4 hello"), "mode": "check"})
        self.assertIn("Excel", r.context["error"])
        r = self.client.post(reverse("reports:data_import", args=["products"]), {"mode": "check"})
        self.assertIn("Choose", r.context["error"])

    def test_csv_upload_also_works(self):
        f = SimpleUploadedFile("p.csv", b"SKU,Name,Selling price\nCSV-1,From csv,9\n", content_type="text/csv")
        res = self.client.post(reverse("reports:data_import", args=["products"]), {"file": f, "mode": "import"}).context["result"]
        self.assertEqual(res["created"], 1)

    # -------------------------------------------------------------- customers etc.
    def test_customers_create_then_update_by_phone(self):
        res = self.imp("customers", ["Name", "Phone", "Email", "Address"], [["Ali Khan", "0300", "a@x.com", "Lahore"]])
        self.assertEqual(res["created"], 1)
        res = self.imp("customers", ["Name", "Phone", "Address"], [["Ali K.", "0300", "Karachi"]])
        self.assertEqual((res["created"], res["updated"]), (0, 1))
        c = Customer.objects.get()
        self.assertEqual((c.name, c.address, c.email), ("Ali K.", "Karachi", "a@x.com"))

    def test_customer_bad_cnic_is_reported(self):
        res = self.imp("customers", ["Name", "CNIC"], [["Bad", "123"], ["Good", "35201-1234567-1"]])
        self.assertEqual((res["created"], res["failed"]), (1, 1))

    def test_suppliers_matched_by_name(self):
        self.imp("suppliers", ["Name", "Phone"], [["Acme", "1"]])
        res = self.imp("suppliers", ["Name", "Phone"], [["acme", "2"], ["Other", "3"]])
        self.assertEqual((res["created"], res["updated"]), (1, 1))
        self.assertEqual(Supplier.objects.get(name="Acme").phone, "2")

    def test_categories_and_brands(self):
        res = self.imp("categories", ["Name", "Parent category"], [["Soft drinks", "Drinks"], ["Soft drinks", "Drinks"]])
        self.assertEqual((res["created"], res["skipped"]), (1, 1))
        self.assertEqual(Category.objects.get(name="Soft drinks").parent.name, "Drinks")
        res = self.imp("brands", ["Name"], [["Lipton"], ["lipton"], ["Tapal"]])
        self.assertEqual((res["created"], res["skipped"]), (2, 1))

    # ----------------------------------------------------------------- expenses
    def test_expenses_import_and_reimport_is_safe(self):
        rows = [["2026-10-01", "Rent", "25,000", "cash", "Landlord", "October", "EXP-IMP-1"], ["2026-10-02", "Tea", 300, "card", "", "", "EXP-IMP-2"]]
        h = ["Date", "Category", "Amount", "Method", "Paid to", "Note", "Reference"]
        res = self.imp("expenses", h, rows)
        self.assertEqual((res["created"], res["failed"]), (2, 0))
        self.assertEqual(Expense.objects.get(reference="EXP-IMP-1").amount, 25000)
        self.assertEqual(ExpenseCategory.objects.filter(store=self.w.store).count(), 2)
        res = self.imp("expenses", h, rows)
        self.assertEqual((res["created"], res["skipped"]), (0, 2))
        self.assertEqual(Expense.objects.count(), 2)

    def test_expense_row_errors(self):
        res = self.imp("expenses", ["Date", "Category", "Amount", "Method"],
                       [["not a date", "Rent", 5, "cash"], ["2026-10-01", "Rent", 0, "cash"], ["2026-10-01", "Rent", 5, "bitcoin"], ["2026-10-01", "Rent", 5, ""]])
        self.assertEqual((res["created"], res["failed"]), (1, 3))
        e = Expense.objects.get()
        self.assertTrue(e.reference)                                            # numbered automatically
        self.assertEqual((e.method, e.created_by_id), ("cash", self.w.manager.pk))

    # ---------------------------------------------------------------- purchases
    def test_purchase_lines_are_grouped_and_received_into_stock(self):
        h = ["Reference", "Date", "Supplier", "Product SKU", "Quantity", "Unit cost"]
        res = self.imp("purchases", h, [["PO-X1", "2026-10-01", "New Supplier", "TEA-1", 10, 20], ["PO-X1", "2026-10-01", "New Supplier", "MILK-1", 5, 30],
                                        ["", "2026-10-02", "New Supplier", "TEA-1", 1, 25]])
        self.assertEqual((res["created"], res["failed"]), (2, 0))
        self.assertEqual(dict(res["made"]), {"suppliers": 1})
        po = PurchaseOrder.objects.get(reference="PO-X1")
        self.assertEqual((po.items.count(), po.total_amount, po.paid_amount, po.status), (2, 350, 0, "received"))
        self.assertEqual(StockItem.objects.get(variant=self.w.tea, warehouse=self.w.warehouse).quantity, 21)        # 10 + 10 + 1
        self.assertEqual(StockItem.objects.get(variant=self.w.milk, warehouse=self.w.warehouse).quantity, 10)
        self.assertTrue(PurchaseOrder.objects.exclude(reference="PO-X1").get().reference)                         # auto-numbered

    def test_purchase_reimport_skips_and_bad_lines_fail_the_whole_purchase(self):
        h = ["Reference", "Date", "Supplier", "Product SKU", "Quantity", "Unit cost", "Status"]
        self.imp("purchases", h, [["PO-1", "2026-10-01", "S", "TEA-1", 1, 1, "ordered"]])
        res = self.imp("purchases", h, [["PO-1", "2026-10-01", "S", "TEA-1", 1, 1, "ordered"], ["PO-2", "2026-10-01", "S", "TEA-1", 1, 1, ""], ["PO-2", "2026-10-01", "S", "GHOST", 1, 1, ""]])
        self.assertEqual((res["created"], res["skipped"], res["failed"]), (0, 1, 1))
        self.assertIn("GHOST", res["errors"][0]["message"])
        self.assertFalse(PurchaseOrder.objects.filter(reference="PO-2").exists())
        self.assertEqual(StockItem.objects.get(variant=self.w.tea, warehouse=self.w.warehouse).quantity, 10)      # 'ordered' adds nothing

    def test_purchase_export_round_trip(self):
        h = ["Reference", "Date", "Supplier", "Product SKU", "Quantity", "Unit cost"]
        self.imp("purchases", h, [["PO-R", "2026-10-01", "S", "TEA-1", 2, 5]])
        rows = sheet_rows(self.client.get(reverse("reports:data_export", args=["purchases"])))
        self.assertEqual(rows[0][:2], ["Reference", "Date"])
        self.assertEqual((rows[1][0], rows[1][5], rows[1][6], rows[1][7]), ("PO-R", "TEA-1", 2, 5))

    # --------------------------------------------------- scope & permissions
    def test_other_branch_data_is_never_exported_or_matched(self):
        other = Store.objects.create(name="Other", company=self.w.company)
        unit = Unit.objects.create(store=other, name="Piece", short_name="pc")
        Product.objects.create(store=other, unit=unit, name="Secret", sku="SECRET")
        Customer.objects.create(store=other, name="Far Away", phone="0999")
        rows = sheet_rows(self.client.get(reverse("reports:data_export", args=["products"])))
        self.assertNotIn("SECRET", [r[0] for r in rows])
        res = self.imp("customers", ["Name", "Phone"], [["Mine", "0999"]])
        self.assertEqual((res["created"], res["updated"]), (1, 0))                 # same phone in another branch is a different customer
        res = self.imp("products", ["SKU", "Name", "Selling price"], [["SECRET", "Mine too", 5]])
        self.assertEqual(res["created"], 1)
        self.assertEqual(Product.objects.filter(sku="SECRET").count(), 2)

    def test_permissions(self):
        self.client.logout()
        self.client.login(username="cashier", password=PASSWORD)                     # sales.view_sale only
        self.assertEqual(self.client.get(reverse("reports:data_hub")).status_code, 200)
        self.assertEqual(self.client.get(reverse("reports:data_export", args=["sales"])).status_code, 200)
        for key in ("products", "customers", "purchases", "expenses", "staff"):
            self.assertEqual(self.client.get(reverse("reports:data_export", args=[key])).status_code, 403, key)
            self.assertEqual(self.client.get(reverse("reports:data_template", args=[key])).status_code, 403, key)
            self.assertEqual(self.client.post(reverse("reports:data_import", args=[key]), {"file": xlsx(["Name"], [["x"]])}).status_code, 403, key)
        hub = self.client.get(reverse("reports:data_hub"))
        self.assertNotContains(hub, "Import from Excel")
        grant(self.w.cashier, "catalog.view_product", "catalog.add_product")
        res = self.client.post(reverse("reports:data_import", args=["products"]), {"file": xlsx(["SKU", "Name", "Selling price"], [["N-1", "New", 5], ["TEA", "", 99]]), "mode": "import"}).context["result"]
        self.assertEqual((res["created"], res["failed"]), (1, 1))                    # may add, may not edit
        self.assertIn("not allowed to edit", res["errors"][0]["message"])
        self.w.tea.refresh_from_db()
        self.assertEqual(self.w.tea.selling_price, 100)

    def test_anonymous_is_sent_to_login(self):
        self.client.logout()
        self.assertEqual(self.client.get(reverse("reports:data_hub")).status_code, 302)
        self.assertEqual(self.client.get(reverse("reports:data_export", args=["products"])).status_code, 302)
