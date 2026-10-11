"""The extra reports (purchases, expenses, registers, suppliers, billers, staff, low stock...) and the Excel / PDF downloads."""
import io
from datetime import timedelta
from decimal import Decimal

from django.test import TestCase
from django.urls import reverse
from django.utils import timezone

from core.pos_fixtures import PASSWORD, PosWorld, grant
from expenses.models import Expense, ExpenseCategory
from inventory.models import StockItem
from payments.models import Payment
from register.models import RegisterSession
from sales.models import Sale, SaleItem
from suppliers.models import PurchaseOrder, PurchaseOrderItem, PurchasePayment, Supplier


class MoreReportTests(TestCase):
    def setUp(self):
        self.w = PosWorld()
        self.today = timezone.localdate()
        self.client.login(username="manager", password=PASSWORD)

    def sale(self, ref, variant, qty, price, biller=None, paid="0"):
        total = Decimal(price) * Decimal(qty)
        s = Sale.objects.create(store=self.w.store, warehouse=self.w.warehouse, reference=ref, biller=biller or self.w.manager, status="completed", subtotal=total, total_amount=total)
        SaleItem.objects.create(sale=s, variant=variant, quantity=qty, unit_price=Decimal(price))
        if Decimal(paid):
            Payment.objects.create(sale=s, reference=f"P-{ref}", amount=Decimal(paid), method="cash", received_by=self.w.manager)
        return s

    def purchase(self, ref, qty, cost, paid="0", status="received", supplier=None, when=None):
        supplier = supplier or Supplier.objects.get_or_create(store=self.w.store, name="Acme Traders", defaults={"phone": "0300"})[0]
        po = PurchaseOrder.objects.create(reference=ref, supplier=supplier, warehouse=self.w.warehouse, status=status, order_date=when or self.today)
        PurchaseOrderItem.objects.create(purchase_order=po, variant=self.w.tea, quantity=qty, cost_price=Decimal(cost))
        if Decimal(paid):
            PurchasePayment.objects.create(reference=f"PP-{ref}", purchase_order=po, amount=Decimal(paid), paid_on=self.today)
        return po

    def data(self, slug, **q):
        r = self.client.get(reverse("reports:pack_report", args=[slug]), q)
        self.assertEqual(r.status_code, 200, slug)
        return r.context["data"]

    def kpis(self, data):
        return {k: v for k, _t, v in data["kpis"]}

    # ------------------------------------------------------------------ sales
    def test_monthly_sales_and_categories(self):
        self.sale("S-1", self.w.tea, 2, "100", paid="200")
        self.sale("S-2", self.w.milk, 3, "50")
        rows = self.data("monthly-sales")["rows"]
        self.assertEqual((len(rows), rows[0]["n"], rows[0]["total"], rows[0]["paid"], rows[0]["due"]), (1, 2, 350, 200, 150))
        cat = self.data("categories")["rows"]
        self.assertEqual((cat[0]["category"], cat[0]["qty"], cat[0]["revenue"]), ("Uncategorised", 5, 350))      # fixture products have no category

    def test_billers_and_staff(self):
        self.sale("S-1", self.w.tea, 1, "100", biller=self.w.cashier, paid="100")
        self.sale("S-2", self.w.tea, 3, "100", biller=self.w.manager)
        billers = {r["biller"]: r for r in self.data("billers")["rows"]}
        self.assertEqual((billers["manager"]["total"], billers["cashier"]["total"], billers["cashier"]["due"]), (300, 100, 0))
        RegisterSession.objects.create(register=self.w.counter1, opened_by=self.w.cashier, opening_cash=0)
        staff = {r["staff"]: r for r in self.data("staff")["rows"]}
        self.assertEqual((staff["cashier"]["n"], staff["cashier"]["shifts"], staff["manager"]["total"]), (1, 1, 300))

    # -------------------------------------------------------------- purchases
    def test_purchases_totals_and_filters(self):
        self.purchase("PO-1", 10, "20", paid="50")                       # 200, paid 50
        self.purchase("PO-2", 5, "10", status="ordered")                 # 50
        self.purchase("PO-3", 1, "999", status="cancelled")              # ignored
        d = self.data("purchases")
        self.assertEqual([(r["reference"], r["total"], r["paid"], r["due"]) for r in d["rows"]], [("PO-2", 50, 0, 50), ("PO-1", 200, 50, 150)])
        k = self.kpis(d)
        self.assertEqual((k["Purchases"], k["Total purchased"], k["Still owed"]), (2, 250, 200))
        self.assertEqual([r["reference"] for r in self.data("purchases", status="received")["rows"]], ["PO-1"])
        other = Supplier.objects.create(store=self.w.store, name="Other")
        self.assertEqual(self.data("purchases", supplier=other.pk)["rows"], [])

    def test_purchase_dates_are_respected_and_grouped(self):
        self.purchase("PO-1", 1, "100")
        self.purchase("PO-2", 2, "100")
        self.purchase("PO-OLD", 1, "500", when=self.today - timedelta(days=90))
        daily = self.data("daily-purchases")["rows"]
        self.assertEqual((len(daily), daily[0]["n"], daily[0]["total"]), (1, 2, 300))
        self.assertEqual(self.data("monthly-purchases", **{"from": str(self.today - timedelta(days=120))})["kpis"][1][2], 800)
        self.assertEqual(self.data("suppliers")["rows"][0]["total"], 300)

    def test_purchase_returns_reduce_what_is_owed(self):
        from suppliers.models import PurchaseReturn, PurchaseReturnItem
        po = self.purchase("PO-1", 10, "20")                                                     # 200
        ret = PurchaseReturn.objects.create(reference="PR-1", purchase_order=po)
        PurchaseReturnItem.objects.create(purchase_return=ret, variant=self.w.tea, quantity=2, cost_price=Decimal("20"))
        row = self.data("purchases")["rows"][0]
        self.assertEqual((row["total"], row["returned"], row["due"]), (200, 40, 160))

    # --------------------------------------------------------------- expenses
    def test_expenses_and_filters(self):
        rent = ExpenseCategory.objects.create(store=self.w.store, name="Rent")
        tea = ExpenseCategory.objects.create(store=self.w.store, name="Tea")
        Expense.objects.create(store=self.w.store, category=rent, reference="E-1", amount=Decimal("100"), expense_date=self.today, method="cash")
        Expense.objects.create(store=self.w.store, category=tea, reference="E-2", amount=Decimal("30"), expense_date=self.today, method="card")
        d = self.data("expenses")
        self.assertEqual((self.kpis(d)["Total expenses"], self.kpis(d)["Entries"]), (130, 2))
        self.assertEqual([r["reference"] for r in self.data("expenses", category=rent.pk)["rows"]], ["E-1"])
        self.assertEqual([r["reference"] for r in self.data("expenses", method="card")["rows"]], ["E-2"])

    # -------------------------------------------------------------- registers
    def test_register_shifts_report(self):
        s = RegisterSession.objects.create(register=self.w.counter1, opened_by=self.w.cashier, opening_cash=Decimal("500"))
        sale = self.sale("S-1", self.w.tea, 1, "100", paid="100")
        sale.register_session = s
        sale.save()
        Payment.objects.filter(sale=sale).update(register_session=s)
        s.closed_by, s.closed_at, s.status, s.closing_cash_counted = self.w.cashier, timezone.now(), "closed", Decimal("590")
        s.expected_cash_at_close = Decimal("600")
        s.save()
        RegisterSession.objects.create(register=self.w.counter2, opened_by=self.w.cashier2, opening_cash=Decimal("0"))
        d = self.data("registers")
        by = {r["register"]: r for r in d["rows"]}
        self.assertEqual((by["Counter 1"]["sales"], by["Counter 1"]["cash_sales"], by["Counter 1"]["expected"], by["Counter 1"]["variance"]), (100, 100, 600, -10))
        self.assertIsNone(by["Counter 2"]["variance"])
        self.assertEqual((self.kpis(d)["Shifts"], self.kpis(d)["Still open"], self.kpis(d)["Net cash difference"]), (2, 1, -10))
        self.assertEqual(len(self.data("registers", status="closed")["rows"]), 1)
        self.assertEqual(len(self.data("registers", register=self.w.counter2.pk)["rows"]), 1)

    # -------------------------------------------------------------- low stock
    def test_low_stock_alerts(self):
        self.w.tea.product.alert_quantity = Decimal("15")             # 10 on hand -> low
        self.w.tea.product.save()
        self.w.milk.product.alert_quantity = Decimal("2")             # 5 on hand -> fine
        self.w.milk.product.save()
        rows = self.data("stock-alerts")["rows"]
        self.assertEqual([(r["sku"], r["qty"], r["alert"], r["short"]) for r in rows], [("EMPTY-1", 0, 0, 0), ("TEA-1", 10, 15, 5)])   # sold-out items are listed too
        self.assertEqual(self.kpis(self.data("stock-alerts"))["Out of stock"], 1)

    # ---------------------------------------------------------------- scoping
    def test_other_branch_data_never_leaks(self):
        from tenants.models import Store, Warehouse
        other = Store.objects.create(name="Other", company=self.w.company)
        ow = Warehouse.objects.create(store=other, name="O")
        sup = Supplier.objects.create(store=other, name="Far")
        po = PurchaseOrder.objects.create(reference="PO-X", supplier=sup, warehouse=ow, status="received", order_date=self.today)
        PurchaseOrderItem.objects.create(purchase_order=po, variant=self.w.tea, quantity=1, cost_price=Decimal("77"))
        self.assertEqual(self.data("purchases")["rows"], [])
        self.assertEqual(self.data("suppliers")["rows"], [])

    # ------------------------------------------------------------ permissions
    def test_permissions_hide_reports(self):
        self.client.logout()
        self.client.login(username="cashier", password=PASSWORD)                     # cashier has sales.view_sale only
        self.assertEqual(self.client.get(reverse("reports:pack_report", args=["monthly-sales"])).status_code, 200)
        for slug in ("purchases", "expenses", "staff", "registers", "stock-alerts"):
            self.assertEqual(self.client.get(reverse("reports:pack_report", args=[slug])).status_code, 403, slug)
            self.assertEqual(self.client.get(reverse("reports:pack_report_export", args=[slug, "xlsx"])).status_code, 403, slug)
        index = self.client.get(reverse("reports:pack_index"))
        self.assertContains(index, "Monthly sales")
        self.assertNotContains(index, "Purchases")
        grant(self.w.cashier, "suppliers.view_purchaseorder")
        self.assertEqual(self.client.get(reverse("reports:pack_report", args=["purchases"])).status_code, 200)

    # ------------------------------------------------------------ downloads
    def test_excel_export_has_the_same_numbers(self):
        import openpyxl
        self.purchase("PO-1", 10, "20", paid="50")
        r = self.client.get(reverse("reports:pack_report_export", args=["purchases", "xlsx"]))
        self.assertEqual(r.status_code, 200)
        self.assertIn("spreadsheetml", r["Content-Type"])
        self.assertIn("attachment", r["Content-Disposition"])
        ws = openpyxl.load_workbook(io.BytesIO(r.content)).active
        cells = [[c for c in row] for row in ws.iter_rows(values_only=True)]
        header = next(i for i, row in enumerate(cells) if row and row[0] == "Reference")
        self.assertEqual(cells[header + 1][0], "PO-1")
        self.assertIn(200, cells[header + 1])
        self.assertIn(150, cells[header + 1])

    def test_excel_cells_cannot_run_formulas(self):
        import openpyxl
        sup = Supplier.objects.create(store=self.w.store, name='=HYPERLINK("http://x","y")')
        self.purchase("PO-1", 1, "5", supplier=sup)
        r = self.client.get(reverse("reports:pack_report_export", args=["purchases", "xlsx"]))
        values = [c for row in openpyxl.load_workbook(io.BytesIO(r.content)).active.iter_rows(values_only=True) for c in row]
        self.assertNotIn('=HYPERLINK("http://x","y")', values)
        self.assertIn("'=HYPERLINK(\"http://x\",\"y\")", values)

    def test_pdf_export_for_every_report(self):
        from reports.pack import REGISTRY
        self.purchase("PO-1", 10, "20", paid="50")
        self.sale("S-1", self.w.tea, 2, "100", paid="200")
        for slug in REGISTRY:
            r = self.client.get(reverse("reports:pack_report_export", args=[slug, "pdf"]))
            self.assertEqual(r.status_code, 200, slug)
            self.assertEqual(r["Content-Type"], "application/pdf", slug)
            self.assertTrue(r.content.startswith(b"%PDF"), slug)
            x = self.client.get(reverse("reports:pack_report_export", args=[slug, "xlsx"]))
            self.assertEqual(x.status_code, 200, slug)

    def test_unknown_format_is_404(self):
        self.assertEqual(self.client.get(reverse("reports:pack_report_export", args=["sales", "exe"])).status_code, 404)
