"""Excel (openpyxl) and optional PDF (WeasyPrint) output for any report."""
import io
from datetime import date, datetime
from decimal import Decimal
from html import escape


def to_xlsx(title, headers, rows, *, subtitle=""):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    ws.append([title])
    ws["A1"].font = Font(bold=True, size=14)
    if subtitle:
        ws.append([subtitle])
    ws.append([])
    ws.append(list(headers))
    header_row = ws.max_row
    for cell in ws[header_row]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="374151")
        cell.alignment = Alignment(horizontal="center")
    for row in rows:
        ws.append([float(v) if isinstance(v, Decimal) else v for v in row])
    for row in ws.iter_rows(min_row=header_row + 1):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "#,##0.00"
            elif isinstance(cell.value, (date, datetime)):
                cell.number_format = "yyyy-mm-dd"
    for col in ws.columns:
        width = max((len(str(c.value)) for c in col if c.value is not None), default=8)
        ws.column_dimensions[col[0].column_letter].width = min(max(width + 2, 10), 50)
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def to_html_table(title, headers, rows, subtitle=""):
    head = "".join(f"<th>{escape(str(h))}</th>" for h in headers)
    body = "".join("<tr>" + "".join(f"<td>{escape(str(v))}</td>" for v in row) + "</tr>" for row in rows)
    return (f"<html><head><meta charset='utf-8'><style>body{{font-family:sans-serif;font-size:11px}}table{{border-collapse:collapse;width:100%}}"
            f"th,td{{border:1px solid #ccc;padding:3px 6px;text-align:left}}th{{background:#374151;color:#fff}}</style></head>"
            f"<body><h2>{escape(title)}</h2><p>{escape(subtitle)}</p><table><thead><tr>{head}</tr></thead><tbody>{body}</tbody></table></body></html>")


def to_pdf(title, headers, rows, subtitle=""):
    """Returns PDF bytes, or None when WeasyPrint (and its system libraries) is not available on this machine."""
    try:
        from weasyprint import HTML
    except Exception:                                    # noqa: BLE001
        return None
    try:
        return HTML(string=to_html_table(title, headers, rows, subtitle)).write_pdf()
    except Exception:                                    # noqa: BLE001
        return None
