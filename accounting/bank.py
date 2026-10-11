"""
Bank statements and reconciliation.

Excel statement columns (header row, any order, case-insensitive):
    date | description | reference | debit | credit        (withdrawals in debit, deposits in credit)
 or date | description | reference | amount                (+ deposit, - withdrawal)

Nothing is imported unless EVERY row is valid (and not a duplicate):
preview_statement() shows the problems, import_statement() refuses to run
when there are any.
"""
import uuid
from datetime import date, datetime, timedelta
from decimal import Decimal, InvalidOperation

from django.db import transaction
from django.db.models import Sum

from core.errors import ERPError
from core.money import ZERO, D, q2
from .models import BankStatementLine, JournalLine, POSTED_STATUSES


class StatementImportError(ERPError):
    def __init__(self, errors):
        self.errors = errors
        super().__init__(f"{len(errors)} problem(s) found in the statement; nothing was imported.")


def _to_date(value):
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = str(value or "").strip()
    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%d/%m/%Y", "%m/%d/%Y"):
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    raise ValueError(f"unreadable date '{value}'")


def _to_decimal(value):
    if value is None or str(value).strip() == "":
        return ZERO
    try:
        return Decimal(str(value).replace(",", "").strip())
    except InvalidOperation:
        raise ValueError(f"unreadable amount '{value}'")


def preview_statement(bank_account, file):
    """Parse an .xlsx file -> (rows, errors). rows are ready-to-save dicts."""
    from openpyxl import load_workbook
    sheet = load_workbook(file, read_only=True, data_only=True).active
    iterator = sheet.iter_rows(values_only=True)
    try:
        header = [str(h or "").strip().lower() for h in next(iterator)]
    except StopIteration:
        return [], [{"row": 1, "error": "The file is empty."}]
    if "date" not in header or not ({"amount"} <= set(header) or {"debit", "credit"} & set(header)):
        return [], [{"row": 1, "error": "Header must contain 'date' and either 'amount' or 'debit'/'credit'."}]

    idx = {name: header.index(name) for name in header if name}
    rows, errors, seen = [], [], set()
    existing = set(BankStatementLine.objects.filter(bank_account=bank_account).values_list("date", "reference", "amount"))
    for number, raw in enumerate(iterator, start=2):
        if raw is None or all(v in (None, "") for v in raw):
            continue
        get = lambda col: raw[idx[col]] if col in idx and idx[col] < len(raw) else None
        try:
            line_date = _to_date(get("date"))
            if "amount" in idx and get("amount") not in (None, ""):
                amount = _to_decimal(get("amount"))
            else:
                amount = _to_decimal(get("credit")) - _to_decimal(get("debit"))
            if amount == 0:
                raise ValueError("amount is zero")
        except ValueError as exc:
            errors.append({"row": number, "error": str(exc)})
            continue
        reference = str(get("reference") or "")[:80]
        key = (line_date, reference, q2(amount))
        if key in existing or key in seen:
            errors.append({"row": number, "error": "duplicate of an already imported/duplicated line"})
            continue
        seen.add(key)
        rows.append({"date": line_date, "description": str(get("description") or "")[:255],
                     "reference": reference, "amount": q2(amount)})
    return rows, errors


@transaction.atomic
def import_statement(bank_account, file):
    rows, errors = preview_statement(bank_account, file)
    if errors:
        raise StatementImportError(errors)
    batch = uuid.uuid4().hex[:12]
    BankStatementLine.objects.bulk_create([BankStatementLine(bank_account=bank_account, import_batch=batch, **r) for r in rows])
    return len(rows), batch


def _unmatched_book_lines(bank_account):
    return JournalLine.objects.filter(
        account=bank_account.ledger_account, entry__status__in=POSTED_STATUSES, statement_line__isnull=True,
    ).select_related("entry")


@transaction.atomic
def auto_match(bank_account, *, tolerance_days=3):
    """Pair each unmatched statement line with one unmatched ledger line of the
    same amount whose entry date is within `tolerance_days`. Returns the number of pairs."""
    matched = 0
    used = set()
    book = list(_unmatched_book_lines(bank_account))
    for stmt in BankStatementLine.objects.filter(bank_account=bank_account, matched_line__isnull=True).order_by("date", "id"):
        best, best_gap = None, None
        for line in book:
            if line.id in used:
                continue
            signed = line.debit - line.credit       # + money in to the bank ledger
            if q2(signed) != q2(stmt.amount):
                continue
            gap = abs((line.entry.date - stmt.date).days)
            if gap <= tolerance_days and (best is None or gap < best_gap):
                best, best_gap = line, gap
        if best is not None:
            stmt.matched_line = best
            stmt.save(update_fields=["matched_line", "updated_at"])
            used.add(best.id)
            matched += 1
    return matched


@transaction.atomic
def match(statement_line, journal_line):
    if journal_line.account_id != statement_line.bank_account.ledger_account_id:
        raise ERPError("That journal line is not on this bank account's ledger.")
    if q2(journal_line.debit - journal_line.credit) != q2(statement_line.amount):
        raise ERPError("Amounts differ; cannot match.")
    if hasattr(journal_line, "statement_line") and journal_line.statement_line is not None:
        raise ERPError("That journal line is already matched.")
    statement_line.matched_line = journal_line
    statement_line.save(update_fields=["matched_line", "updated_at"])


def unmatch(statement_line):
    statement_line.matched_line = None
    statement_line.save(update_fields=["matched_line", "updated_at"])


def reconciliation_report(bank_account, *, as_of=None):
    as_of = as_of or date.today()
    stmt_qs = BankStatementLine.objects.filter(bank_account=bank_account, date__lte=as_of)
    statement_balance = bank_account.opening_statement_balance + (stmt_qs.aggregate(t=Sum("amount"))["t"] or ZERO)
    book_agg = JournalLine.objects.filter(
        account=bank_account.ledger_account, entry__status__in=POSTED_STATUSES, entry__date__lte=as_of
    ).aggregate(d=Sum("debit"), c=Sum("credit"))
    book_balance = (book_agg["d"] or ZERO) - (book_agg["c"] or ZERO)
    unmatched_statement = list(stmt_qs.filter(matched_line__isnull=True))
    unmatched_book = [l for l in _unmatched_book_lines(bank_account) if l.entry.date <= as_of]
    return {
        "statement_balance": q2(statement_balance), "book_balance": q2(book_balance),
        "difference": q2(statement_balance - book_balance),
        "unmatched_statement_lines": unmatched_statement, "unmatched_book_lines": unmatched_book,
    }
